import io

import docx
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.db.base import Base
from app.db.session import get_db
from app.main import app

_engine = create_engine("sqlite:///./test_documents.db", connect_args={"check_same_thread": False})
_TestSessionLocal = sessionmaker(bind=_engine, autoflush=False, autocommit=False)


def _override_get_db():
    db = _TestSessionLocal()
    try:
        yield db
    finally:
        db.close()


@pytest.fixture(autouse=True, scope="module")
def _setup_database():
    Base.metadata.create_all(bind=_engine)
    app.dependency_overrides[get_db] = _override_get_db
    yield
    app.dependency_overrides.clear()
    Base.metadata.drop_all(bind=_engine)


client = TestClient(app)


def _build_sample_docx() -> bytes:
    document = docx.Document()
    document.add_heading("Non-Disclosure Agreement", level=1)
    document.add_paragraph("This Agreement is entered into between Client and Vendor.")

    bullet = document.add_paragraph("Confidential information must not be disclosed.")
    bullet.style = document.styles["List Bullet"]

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Term"
    table.cell(0, 1).text = "Definition"
    table.cell(1, 0).text = "Effective Date"
    table.cell(1, 1).text = "The date of signing"

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_upload_valid_docx_returns_summary() -> None:
    content = _build_sample_docx()

    response = client.post(
        "/api/v1/documents",
        files={
            "file": (
                "sample.docx",
                content,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 201
    body = response.json()
    assert body["status"] == "processed"
    assert body["paragraph_count"] > 0
    assert body["table_count"] == 1
    assert body["section_count"] == 1


def test_upload_rejects_non_docx_extension() -> None:
    response = client.post(
        "/api/v1/documents",
        files={"file": ("notes.txt", b"just some text", "text/plain")},
    )

    assert response.status_code == 415
    assert response.json()["error"]["code"] == "unsupported_file_type"


def test_upload_rejects_corrupt_docx() -> None:
    response = client.post(
        "/api/v1/documents",
        files={
            "file": (
                "broken.docx",
                b"not a real docx file",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "corrupt_document"


def test_retrieve_document_and_structure_with_stable_ids() -> None:
    content = _build_sample_docx()
    upload_response = client.post(
        "/api/v1/documents",
        files={
            "file": (
                "sample2.docx",
                content,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    document_id = upload_response.json()["id"]

    get_response = client.get(f"/api/v1/documents/{document_id}")
    assert get_response.status_code == 200
    assert get_response.json()["id"] == document_id

    structure_response = client.get(f"/api/v1/documents/{document_id}/structure")
    assert structure_response.status_code == 200
    structure = structure_response.json()

    assert structure["document"]["id"] == document_id
    assert len(structure["sections"]) == 1
    assert len(structure["tables"]) == 1
    assert len(structure["paragraphs"]) > 0

    paragraph_ids = [p["id"] for p in structure["paragraphs"]]
    assert len(paragraph_ids) == len(set(paragraph_ids))

    heading = next(p for p in structure["paragraphs"] if p["paragraph_type"] == "heading")
    assert heading["heading_level"] == 1

    list_item = next(p for p in structure["paragraphs"] if p["paragraph_type"] == "list_item")
    assert list_item["list_level"] == 0

    table_cells = [p for p in structure["paragraphs"] if p["location"] == "table_cell"]
    assert len(table_cells) == 4


def test_retrieve_unknown_document_returns_404() -> None:
    response = client.get("/api/v1/documents/does-not-exist")
    assert response.status_code == 404


def test_list_documents_carries_review_status() -> None:
    content = _build_sample_docx()
    upload_response = client.post(
        "/api/v1/documents",
        files={
            "file": (
                "sample3.docx",
                content,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    document_id = upload_response.json()["id"]

    list_response = client.get("/api/v1/documents")
    assert list_response.status_code == 200
    items = list_response.json()

    # SQLite's created_at only has second resolution, so exact ordering
    # against documents uploaded by earlier tests in this module isn't
    # reliable here (Postgres uses microsecond precision in practice) —
    # just confirm this document is present with correctly joined fields.
    match = next(item for item in items if item["id"] == document_id)
    assert match["agreement_type"] is None
    assert match["latest_review_id"] is None
    assert match["review_status"] is None
    assert match["has_reviewed_document"] is False
