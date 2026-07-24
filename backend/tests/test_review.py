import io
from dataclasses import dataclass, field

import docx
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.core.exceptions import (
    AIRefusalError,
    AIResponseInvalidError,
    ClaudeRefusalError,
    ClaudeResponseInvalidError,
)
from app.db.base import Base
from app.db.session import get_db
from app.main import app
from app.models.document import DocumentParagraph
from app.services import agreement_detector, party_normalization_service, review_engine
from app.services.claude_service import ClaudeReviewResult, ClaudeService
from app.services.gemini_service import GeminiReviewResult, GeminiService, _to_gemini_schema
from app.services.prompt_builder import FINDINGS_JSON_SCHEMA
from app.services.review_validator import validate_findings

_engine = create_engine("sqlite:///./test_review.db", connect_args={"check_same_thread": False})
_TestSessionLocal = sessionmaker(bind=_engine, autoflush=False, autocommit=False)


def _override_get_db():
    db = _TestSessionLocal()
    try:
        yield db
    finally:
        db.close()


@pytest.fixture(autouse=True, scope="module")
def _setup_database():
    Base.metadata.create_all(bind=_engine)
    app.dependency_overrides[get_db] = _override_get_db
    yield
    app.dependency_overrides.clear()
    Base.metadata.drop_all(bind=_engine)


client = TestClient(app)


# ---------------------------------------------------------------------------
# Agreement detection
# ---------------------------------------------------------------------------


def test_agreement_detector_uses_heading_first():
    paragraphs = [
        DocumentParagraph(
            id="p1", location="body", paragraph_type="heading", order_index=0,
            text="Non-Disclosure Agreement",
        ),
        DocumentParagraph(
            id="p2", location="body", paragraph_type="body", order_index=1,
            text="This agreement covers confidential information.",
        ),
    ]
    assert agreement_detector.detect_agreement_type(paragraphs) == "NDA"


def test_agreement_detector_falls_back_to_keywords_when_no_heading_match():
    paragraphs = [
        DocumentParagraph(
            id="p1", location="body", paragraph_type="body", order_index=0,
            text="This is a master service agreement governing the statement of work process.",
        ),
    ]
    assert agreement_detector.detect_agreement_type(paragraphs) == "MSA"


def test_agreement_detector_returns_unknown_for_empty_document():
    assert agreement_detector.detect_agreement_type([]) == "Unknown"


# ---------------------------------------------------------------------------
# Party normalization
# ---------------------------------------------------------------------------


def test_party_normalization_replaces_defined_companies_with_roles():
    paragraphs = [
        DocumentParagraph(
            id="p1", location="body", paragraph_type="body", order_index=0,
            text='This Agreement is between ABC Technologies Pvt Ltd ("Supplier") '
            'and XYZ Corporation ("Client").',
        ),
        DocumentParagraph(
            id="p2", location="body", paragraph_type="body", order_index=1,
            text="ABC Technologies Pvt Ltd shall not disclose XYZ Corporation's confidential information.",
        ),
    ]

    result = party_normalization_service.normalize_parties(paragraphs)

    names = {entry.original_name for entry in result.mapping}
    assert "ABC Technologies Pvt Ltd" in names
    assert "XYZ Corporation" in names

    normalized_p2 = result.normalized_text_by_paragraph_id["p2"]
    assert "ABC Technologies Pvt Ltd" not in normalized_p2
    assert "XYZ Corporation" not in normalized_p2
    assert "Supplier" in normalized_p2
    assert "Client" in normalized_p2


def test_party_normalization_mapping_is_reversible():
    paragraphs = [
        DocumentParagraph(
            id="p1", location="body", paragraph_type="body", order_index=0,
            text='Acme Solutions Inc ("Vendor") will provide services to the Customer.',
        ),
    ]

    result = party_normalization_service.normalize_parties(paragraphs)

    assert len(result.mapping) == 1
    entry = result.mapping[0]
    role_to_name = {entry.normalized_role: entry.original_name}
    assert role_to_name[entry.normalized_role] == "Acme Solutions Inc"


def test_party_normalization_no_op_when_no_companies_present():
    paragraphs = [
        DocumentParagraph(
            id="p1", location="body", paragraph_type="body", order_index=0,
            text="This clause contains no company names at all.",
        ),
    ]
    result = party_normalization_service.normalize_parties(paragraphs)
    assert result.mapping == []
    assert result.normalized_text_by_paragraph_id["p1"] == paragraphs[0].text


# ---------------------------------------------------------------------------
# Claude service (network boundary stubbed out)
# ---------------------------------------------------------------------------


@dataclass
class _FakeTextBlock:
    text: str
    type: str = "text"


@dataclass
class _FakeUsage:
    input_tokens: int = 10
    output_tokens: int = 20


@dataclass
class _FakeResponse:
    content: list = field(default_factory=list)
    stop_reason: str = "end_turn"
    model: str = "claude-opus-4-8"
    usage: _FakeUsage = field(default_factory=_FakeUsage)


class _FakeMessages:
    def __init__(self, response: _FakeResponse) -> None:
        self._response = response

    def create(self, **kwargs):
        return self._response


class _FakeAnthropicClient:
    def __init__(self, response: _FakeResponse) -> None:
        self.messages = _FakeMessages(response)


def test_claude_service_parses_valid_structured_response():
    response = _FakeResponse(
        content=[_FakeTextBlock(text='{"findings": [{"finding_id": "f1", "paragraph_id": "p1", '
                                      '"issue_type": "liability", "severity": "high", '
                                      '"explanation": "one-sided", "suggested_text": "text", '
                                      '"confidence": 0.9}]}')],
    )
    service = ClaudeService(client=_FakeAnthropicClient(response))

    result = service.review_clauses(system="sys", user_message="user", json_schema={})

    assert isinstance(result, ClaudeReviewResult)
    assert result.model == "claude-opus-4-8"
    assert result.findings[0]["finding_id"] == "f1"
    assert result.input_tokens == 10
    assert result.output_tokens == 20


def test_claude_service_raises_on_refusal():
    response = _FakeResponse(content=[], stop_reason="refusal")
    service = ClaudeService(client=_FakeAnthropicClient(response))

    with pytest.raises(ClaudeRefusalError):
        service.review_clauses(system="sys", user_message="user", json_schema={})


def test_claude_service_raises_on_invalid_json():
    response = _FakeResponse(content=[_FakeTextBlock(text="not json")])
    service = ClaudeService(client=_FakeAnthropicClient(response))

    with pytest.raises(ClaudeResponseInvalidError):
        service.review_clauses(system="sys", user_message="user", json_schema={})


def test_claude_service_raises_on_truncated_response():
    response = _FakeResponse(content=[_FakeTextBlock(text="{}")], stop_reason="max_tokens")
    service = ClaudeService(client=_FakeAnthropicClient(response))

    with pytest.raises(ClaudeResponseInvalidError):
        service.review_clauses(system="sys", user_message="user", json_schema={})


# ---------------------------------------------------------------------------
# Gemini service (free-tier alternative, same interface as ClaudeService)
# ---------------------------------------------------------------------------


@dataclass
class _FakeGeminiUsage:
    prompt_token_count: int = 10
    candidates_token_count: int = 20


@dataclass
class _FakeGeminiCandidate:
    finish_reason: str = "STOP"


@dataclass
class _FakeGeminiResponse:
    text: str | None = "{}"
    candidates: list = field(default_factory=lambda: [_FakeGeminiCandidate()])
    usage_metadata: _FakeGeminiUsage = field(default_factory=_FakeGeminiUsage)


class _FakeGeminiModels:
    def __init__(self, response: _FakeGeminiResponse) -> None:
        self._response = response

    def generate_content(self, **kwargs):
        return self._response


class _FakeGenaiClient:
    def __init__(self, response: _FakeGeminiResponse) -> None:
        self.models = _FakeGeminiModels(response)


def test_gemini_service_parses_valid_structured_response():
    response = _FakeGeminiResponse(
        text='{"findings": [{"finding_id": "f1", "paragraph_id": "p1", '
        '"issue_type": "liability", "severity": "high", "explanation": "one-sided", '
        '"suggested_text": "text", "confidence": 0.9}]}'
    )
    service = GeminiService(client=_FakeGenaiClient(response))

    result = service.review_clauses(system="sys", user_message="user", json_schema=FINDINGS_JSON_SCHEMA)

    assert isinstance(result, GeminiReviewResult)
    assert result.findings[0]["finding_id"] == "f1"
    assert result.input_tokens == 10
    assert result.output_tokens == 20


def test_gemini_service_raises_on_refusal():
    response = _FakeGeminiResponse(text=None, candidates=[_FakeGeminiCandidate(finish_reason="SAFETY")])
    service = GeminiService(client=_FakeGenaiClient(response))

    with pytest.raises(AIRefusalError):
        service.review_clauses(system="sys", user_message="user", json_schema={})


def test_gemini_service_raises_on_invalid_json():
    response = _FakeGeminiResponse(text="not json")
    service = GeminiService(client=_FakeGenaiClient(response))

    with pytest.raises(AIResponseInvalidError):
        service.review_clauses(system="sys", user_message="user", json_schema={})


def test_gemini_service_raises_on_truncated_response():
    response = _FakeGeminiResponse(text="{}", candidates=[_FakeGeminiCandidate(finish_reason="MAX_TOKENS")])
    service = GeminiService(client=_FakeGenaiClient(response))

    with pytest.raises(AIResponseInvalidError):
        service.review_clauses(system="sys", user_message="user", json_schema={})


def test_gemini_schema_translation_strips_unsupported_keys():
    translated = _to_gemini_schema(FINDINGS_JSON_SCHEMA)

    findings_items = translated["properties"]["findings"]["items"]
    assert "additionalProperties" not in translated
    assert "additionalProperties" not in findings_items
    assert findings_items["properties"]["suggested_text"]["type"] == "string"


def test_resolve_default_service_honors_ai_provider_setting(monkeypatch):
    monkeypatch.setattr(review_engine.get_settings(), "AI_PROVIDER", "claude")
    assert review_engine._resolve_default_service() is review_engine.default_claude_service

    monkeypatch.setattr(review_engine.get_settings(), "AI_PROVIDER", "gemini")
    assert review_engine._resolve_default_service() is review_engine.default_gemini_service


# ---------------------------------------------------------------------------
# Review validator
# ---------------------------------------------------------------------------


def test_review_validator_keeps_valid_and_drops_invalid_findings():
    raw_findings = [
        {
            "finding_id": "f1", "paragraph_id": "p1", "issue_type": "liability",
            "severity": "high", "explanation": "risky", "suggested_text": "fix",
            "confidence": 0.8,
        },
        {  # unknown paragraph_id
            "finding_id": "f2", "paragraph_id": "does-not-exist", "issue_type": "x",
            "severity": "low", "explanation": "x", "suggested_text": None, "confidence": 0.5,
        },
        {  # invalid severity
            "finding_id": "f3", "paragraph_id": "p1", "issue_type": "x",
            "severity": "catastrophic", "explanation": "x", "suggested_text": None, "confidence": 0.5,
        },
        {  # confidence out of range
            "finding_id": "f4", "paragraph_id": "p1", "issue_type": "x",
            "severity": "low", "explanation": "x", "suggested_text": None, "confidence": 1.5,
        },
        {  # missing explanation
            "finding_id": "f5", "paragraph_id": "p1", "issue_type": "x",
            "severity": "low", "explanation": "", "suggested_text": None, "confidence": 0.5,
        },
        {  # duplicate of f1
            "finding_id": "f1", "paragraph_id": "p1", "issue_type": "dup",
            "severity": "low", "explanation": "dup", "suggested_text": None, "confidence": 0.5,
        },
    ]

    validated = validate_findings(raw_findings, valid_paragraph_ids={"p1"})

    assert len(validated) == 1
    assert validated[0].finding_id == "f1"


# ---------------------------------------------------------------------------
# Review API (end-to-end with a stubbed Claude service)
# ---------------------------------------------------------------------------


def _build_sample_docx() -> bytes:
    document = docx.Document()
    document.add_heading("Non-Disclosure Agreement", level=1)
    document.add_paragraph(
        'This Agreement is entered into between ABC Technologies Pvt Ltd ("Supplier") '
        'and XYZ Corporation ("Client").'
    )
    document.add_paragraph(
        "ABC Technologies Pvt Ltd may terminate this Agreement at any time without notice "
        "for any reason, at its sole discretion."
    )
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_review_api_creates_and_retrieves_review(monkeypatch):
    upload_response = client.post(
        "/api/v1/documents",
        files={
            "file": (
                "nda.docx",
                _build_sample_docx(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    document_id = upload_response.json()["id"]

    structure = client.get(f"/api/v1/documents/{document_id}/structure").json()
    target_paragraph = next(
        p for p in structure["paragraphs"]
        if p["location"] == "body" and p["paragraph_type"] != "heading" and "terminate" in p["text"]
    )

    fake_response = _FakeResponse(
        content=[
            _FakeTextBlock(
                text=(
                    '{"findings": [{"finding_id": "f1", '
                    f'"paragraph_id": "{target_paragraph["id"]}", '
                    '"issue_type": "termination_imbalance", "severity": "high", '
                    '"explanation": "One-sided termination right with no notice.", '
                    '"suggested_text": "Either party may terminate with 30 days notice.", '
                    '"confidence": 0.85}]}'
                )
            )
        ],
    )
    fake_service = ClaudeService(client=_FakeAnthropicClient(fake_response))
    monkeypatch.setattr(review_engine, "_resolve_default_service", lambda: fake_service)

    review_response = client.post(f"/api/v1/documents/{document_id}/review")
    assert review_response.status_code == 201
    review_body = review_response.json()

    assert review_body["agreement_type"] == "NDA"
    assert review_body["document_id"] == document_id
    assert len(review_body["findings"]) == 1
    assert review_body["findings"][0]["paragraph_id"] == target_paragraph["id"]
    assert review_body["findings"][0]["severity"] == "high"

    role_names = {entry["normalized_role"] for entry in review_body["party_mapping"]}
    assert "Supplier" in role_names or "Client" in role_names

    get_response = client.get(f"/api/v1/reviews/{review_body['id']}")
    assert get_response.status_code == 200
    assert get_response.json()["id"] == review_body["id"]


def test_review_api_returns_404_for_unknown_review():
    response = client.get("/api/v1/reviews/does-not-exist")
    assert response.status_code == 404


def test_list_reviews_for_document(monkeypatch):
    upload_response = client.post(
        "/api/v1/documents",
        files={
            "file": (
                "nda2.docx",
                _build_sample_docx(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    document_id = upload_response.json()["id"]

    structure = client.get(f"/api/v1/documents/{document_id}/structure").json()
    target_paragraph = next(
        p for p in structure["paragraphs"]
        if p["location"] == "body" and p["paragraph_type"] != "heading" and "terminate" in p["text"]
    )

    fake_response = _FakeResponse(
        content=[
            _FakeTextBlock(
                text=(
                    '{"findings": [{"finding_id": "f1", '
                    f'"paragraph_id": "{target_paragraph["id"]}", '
                    '"issue_type": "termination_imbalance", "severity": "high", '
                    '"explanation": "One-sided.", "suggested_text": "fix", "confidence": 0.8}]}'
                )
            )
        ],
    )
    monkeypatch.setattr(
        review_engine,
        "_resolve_default_service",
        lambda: ClaudeService(client=_FakeAnthropicClient(fake_response)),
    )

    review_response = client.post(f"/api/v1/documents/{document_id}/review")
    review_id = review_response.json()["id"]

    list_response = client.get(f"/api/v1/documents/{document_id}/reviews")
    assert list_response.status_code == 200
    reviews = list_response.json()
    assert len(reviews) == 1
    assert reviews[0]["id"] == review_id
    assert len(reviews[0]["findings"]) == 1


def test_list_reviews_for_unknown_document_returns_404():
    response = client.get("/api/v1/documents/does-not-exist/reviews")
    assert response.status_code == 404
