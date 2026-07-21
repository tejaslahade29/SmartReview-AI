"""Pure DOCX parsing: bytes in, canonical document dataclasses out.

No database access and no file I/O here — this module only understands
DOCX structure. Persistence lives in document_service.py.
"""

import io
import zipfile
from dataclasses import dataclass, field
from enum import Enum

from docx import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.ns import qn
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph

from app.core.exceptions import CorruptDocumentError


class ParagraphLocation(str, Enum):
    BODY = "body"
    HEADER = "header"
    FOOTER = "footer"
    TABLE_CELL = "table_cell"


class ParagraphType(str, Enum):
    HEADING = "heading"
    LIST_ITEM = "list_item"
    BODY = "body"


@dataclass
class ParsedRun:
    text: str
    bold: bool
    italic: bool
    underline: bool
    font_name: str | None
    font_size_pt: float | None


@dataclass
class ParsedParagraph:
    order_index: int
    location: ParagraphLocation
    paragraph_type: ParagraphType
    text: str
    style_name: str | None
    heading_level: int | None
    list_level: int | None
    runs: list[ParsedRun]
    section_index: int | None = None
    table_index: int | None = None
    row_index: int | None = None
    col_index: int | None = None


@dataclass
class ParsedTable:
    order_index: int
    section_index: int | None
    row_count: int
    col_count: int


@dataclass
class ParsedSection:
    order_index: int
    start_type: str


@dataclass
class ParsedDocument:
    sections: list[ParsedSection] = field(default_factory=list)
    tables: list[ParsedTable] = field(default_factory=list)
    paragraphs: list[ParsedParagraph] = field(default_factory=list)

    @property
    def section_count(self) -> int:
        return len(self.sections)

    @property
    def table_count(self) -> int:
        return len(self.tables)

    @property
    def paragraph_count(self) -> int:
        return len(self.paragraphs)

    @property
    def word_count(self) -> int:
        return sum(len(p.text.split()) for p in self.paragraphs)


def _extract_runs(paragraph: DocxParagraph) -> list[ParsedRun]:
    runs: list[ParsedRun] = []
    for run in paragraph.runs:
        font = run.font
        runs.append(
            ParsedRun(
                text=run.text,
                bold=bool(run.bold),
                italic=bool(run.italic),
                underline=bool(run.underline),
                font_name=font.name,
                font_size_pt=font.size.pt if font.size is not None else None,
            )
        )
    return runs


def _classify_paragraph(paragraph: DocxParagraph) -> tuple[ParagraphType, int | None, int | None]:
    style_name = (paragraph.style.name or "") if paragraph.style else ""
    lowered = style_name.lower()

    p_pr = paragraph._p.pPr
    num_pr = p_pr.find(qn("w:numPr")) if p_pr is not None else None

    if lowered.startswith("heading") or lowered == "title":
        digits = "".join(ch for ch in style_name if ch.isdigit())
        heading_level = int(digits) if digits else 0
        return ParagraphType.HEADING, heading_level, None

    if num_pr is not None:
        ilvl_el = num_pr.find(qn("w:ilvl"))
        list_level = int(ilvl_el.get(qn("w:val"))) if ilvl_el is not None else 0
        return ParagraphType.LIST_ITEM, None, list_level

    if lowered.startswith("list"):
        return ParagraphType.LIST_ITEM, None, 0

    return ParagraphType.BODY, None, None


def _has_section_break(paragraph: DocxParagraph) -> bool:
    p_pr = paragraph._p.pPr
    if p_pr is None:
        return False
    return p_pr.find(qn("w:sectPr")) is not None


def _build_paragraph(
    paragraph: DocxParagraph,
    order_index: int,
    location: ParagraphLocation,
    section_index: int | None = None,
    table_index: int | None = None,
    row_index: int | None = None,
    col_index: int | None = None,
) -> ParsedParagraph:
    paragraph_type, heading_level, list_level = _classify_paragraph(paragraph)
    return ParsedParagraph(
        order_index=order_index,
        location=location,
        paragraph_type=paragraph_type,
        text=paragraph.text,
        style_name=paragraph.style.name if paragraph.style else None,
        heading_level=heading_level,
        list_level=list_level,
        runs=_extract_runs(paragraph),
        section_index=section_index,
        table_index=table_index,
        row_index=row_index,
        col_index=col_index,
    )


def _parse_table(
    table: DocxTable,
    order_index: int,
    section_index: int | None,
    paragraph_order_start: int,
) -> tuple[ParsedTable, list[ParsedParagraph]]:
    paragraphs: list[ParsedParagraph] = []
    seen_cell_ids: set[int] = set()
    order_index_counter = paragraph_order_start

    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell_id = id(cell._tc)
            if cell_id in seen_cell_ids:
                continue
            seen_cell_ids.add(cell_id)

            for cell_paragraph in cell.paragraphs:
                paragraphs.append(
                    _build_paragraph(
                        cell_paragraph,
                        order_index=order_index_counter,
                        location=ParagraphLocation.TABLE_CELL,
                        section_index=section_index,
                        table_index=order_index,
                        row_index=row_index,
                        col_index=col_index,
                    )
                )
                order_index_counter += 1

    parsed_table = ParsedTable(
        order_index=order_index,
        section_index=section_index,
        row_count=len(table.rows),
        col_count=len(table.columns),
    )
    return parsed_table, paragraphs


def _parse_header_footer(
    part, location: ParagraphLocation, section_index: int
) -> list[ParsedParagraph]:
    if part is None or part.is_linked_to_previous:
        return []
    paragraphs: list[ParsedParagraph] = []
    for order_index, paragraph in enumerate(part.paragraphs):
        if not paragraph.text.strip() and not paragraph.runs:
            continue
        paragraphs.append(
            _build_paragraph(
                paragraph,
                order_index=order_index,
                location=location,
                section_index=section_index,
            )
        )
    return paragraphs


def parse_docx(content: bytes) -> ParsedDocument:
    try:
        docx_document = DocxDocument(io.BytesIO(content))
    except (PackageNotFoundError, KeyError, ValueError, zipfile.BadZipFile) as exc:
        raise CorruptDocumentError("The uploaded file is not a valid DOCX document.") from exc

    parsed = ParsedDocument()

    for section_index, section in enumerate(docx_document.sections):
        parsed.sections.append(
            ParsedSection(order_index=section_index, start_type=str(section.start_type))
        )
        parsed.paragraphs.extend(
            _parse_header_footer(section.header, ParagraphLocation.HEADER, section_index)
        )
        parsed.paragraphs.extend(
            _parse_header_footer(section.footer, ParagraphLocation.FOOTER, section_index)
        )

    current_section_index = 0
    block_index = 0
    body = docx_document.element.body

    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = DocxParagraph(child, docx_document)
            parsed.paragraphs.append(
                _build_paragraph(
                    paragraph,
                    order_index=block_index,
                    location=ParagraphLocation.BODY,
                    section_index=current_section_index,
                )
            )
            block_index += 1
            if _has_section_break(paragraph) and current_section_index < len(parsed.sections) - 1:
                current_section_index += 1
        elif child.tag == qn("w:tbl"):
            table = DocxTable(child, docx_document)
            parsed_table, table_paragraphs = _parse_table(
                table,
                order_index=block_index,
                section_index=current_section_index,
                paragraph_order_start=0,
            )
            parsed.tables.append(parsed_table)
            parsed.paragraphs.extend(table_paragraphs)
            block_index += 1

    return parsed
